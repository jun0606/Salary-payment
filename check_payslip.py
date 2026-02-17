import openpyxl

wb = openpyxl.load_workbook('test_payslip_with_allowances.xlsx', data_only=True)
ws = wb['오수민']

print('=== 오수민 급여명세서 세부 내역 ===')
print(f'기본급 (D9): {ws["D9"].value:,}원')
print(f'연장근로수당 (D11): {ws["D11"].value:,}원')
print(f'야간근로수당 (D12): {ws["D12"].value:,}원')
print(f'휴일근로수당 (D13): {ws["D13"].value}원')
print(f'주휴수당 (D14): {ws["D14"].value:,}원')
print(f'직급수당 (D15): {ws["D15"].value}원')
print(f'기타수당 (D16): {ws["D16"].value:,}원')
print(f'총 지급액 (D17): {ws["D17"].value:,}원')

print(f'국민연금 (H9): {ws["H9"].value}원')
print(f'건강보험 (H10): {ws["H10"].value}원')
print(f'고용보험 (H11): {ws["H11"].value}원')
print(f'장기요양보험 (H12): {ws["H12"].value}원')
print(f'소득세 (H13): {ws["H13"].value:,}원')
print(f'지방소득세 (H14): {ws["H14"].value:,}원')
print(f'공제합계 (H17): {ws["H17"].value:,}원')

print(f'실지급액 (D18): {ws["D18"].value:,}원')

wb.close()

# 수동 계산 검증
base_pay = 660000
overtime = 15045
night_pay = 0
weekly_allowance = 70210
position_allowance = 0
other_allowance = 500000

total_payment = base_pay + overtime + night_pay + weekly_allowance + position_allowance + other_allowance
print(f'\n수동 계산 총 지급액: {total_payment:,}원')

# 공제 합계 계산
national_pension = 0
health_insurance = 0
employment_insurance = 0
long_term_care = 0
income_tax = 172350
local_income_tax = 17230

total_deduction = national_pension + health_insurance + employment_insurance + long_term_care + income_tax + local_income_tax
net_pay = total_payment - total_deduction

print(f'수동 계산 공제합계: {total_deduction:,}원')
print(f'수동 계산 실지급액: {net_pay:,}원')