import openpyxl

wb = openpyxl.load_workbook('test_payslip_with_allowances.xlsx', data_only=True)
ws = wb['오수민']

print('C열 항목명:')
for i in range(10, 18):
    cell_value = ws[f'C{i}'].value
    print(f'C{i}: {cell_value}')

print('\nD열 금액:')
for i in range(10, 18):
    cell_value = ws[f'D{i}'].value
    print(f'D{i}: {cell_value}')

wb.close()