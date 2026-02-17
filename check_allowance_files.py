import openpyxl
import os

files = ['allowance_files/오수민_교통비_202612.xlsx', 'allowance_files/오수민_식대_202612.xlsx', 'allowance_files/오수민_부장 수당_202612.xlsx']

for file_path in files:
    if os.path.exists(file_path):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        print(f'\n=== {os.path.basename(file_path)} ===')
        print(f'제목: {ws["A1"].value}')
        print(f'직원명: {ws["B3"].value}')
        print(f'수당명: {ws["B4"].value}')
        print(f'수당 유형: {ws["B5"].value}')
        print(f'금액: {ws["B6"].value}')
        print(f'과세 여부: {ws["B7"].value}')

        wb.close()